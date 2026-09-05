"""Build the submission document (DOCX, then PDF via Word) for the application form.

    python scripts/build_submission.py --space https://huggingface.co/spaces/<user>/auto-experimenter

Reads the saved runs under results/ and the screenshots under submission/screenshots/.
"""

from __future__ import annotations

import argparse
import json
import subprocess
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "submission"
SHOTS = OUT_DIR / "screenshots"
AUTHOR = "Aiman El Asad"
REPO = "https://github.com/aimanelasad/auto-experimenter"
KAGGLE = "https://www.kaggle.com/datasets/blastchar/telco-customer-churn"

SUMMARY = (
    "Challenge 2. I built an autonomous experiment loop for churn prediction on the IBM Telco dataset. A planner "
    "reads the ledger of all experiments and proposes the next ones as testable hypotheses; it is either a "
    "deterministic rule set or Claude, whose proposals are validated by code before running. Every experiment is "
    "scored under one fixed 5-fold protocol with identical folds, a judge marks each hypothesis confirmed or "
    "refuted against the incumbent, and only the winner touches the holdout, once. A narrator turns the ledger "
    "into a conclusion, next experiments and retention actions. The Streamlit demo reruns the loop live."
)

CAPTIONS = {
    "01_overview.png": "The demo on load: winner, CV and holdout numbers, every experiment with fold noise, best-so-far line.",
    "02_trail.png": "The hypothesis trail: the planner's rationale per round, each hypothesis with its verdict.",
    "03_live_run.png": "A live run in progress: the loop plans, runs and judges experiments in the browser.",
    "04_narration.png": "The narrator's conclusion and the retention table for the top decile.",
}


def load_run(name: str) -> dict | None:
    d = ROOT / "results" / name
    if not (d / "final.json").exists():
        return None
    final = json.loads((d / "final.json").read_text(encoding="utf-8"))
    run = json.loads((d / "run.json").read_text(encoding="utf-8"))
    lb = pd.read_csv(d / "leaderboard.csv")
    return {"final": final, "run": run, "leaderboard": lb}


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--space", default="https://huggingface.co/spaces/aimanelasad/auto-experimenter")
    parser.add_argument("--no-pdf", action="store_true")
    args = parser.parse_args()
    OUT_DIR.mkdir(exist_ok=True)

    words = len(SUMMARY.split())
    doc = Document()
    doc.core_properties.author = AUTHOR
    doc.core_properties.last_modified_by = AUTHOR
    doc.core_properties.title = "AI challenge submission: Autonomous Experimenter"
    doc.core_properties.subject = "auto-experimenter"
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(1.8)

    doc.add_heading("AI challenge submission: Autonomous Experimenter", level=1)
    doc.add_paragraph(f"{AUTHOR}, {date.today().strftime('%d %B %Y')}. Challenge 2 of the optional AI challenge.")

    doc.add_heading("Links", level=2)
    for label, url in [("Prototype (Hugging Face Space)", args.space), ("Code, results and design notes (GitHub)", REPO),
                       ("Dataset (Kaggle, IBM Telco Customer Churn)", KAGGLE)]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{label}: ").bold = True
        p.add_run(url)

    doc.add_heading(f"Summary ({words} words)", level=2)
    doc.add_paragraph(SUMMARY)

    doc.add_heading("What the prototype does", level=2)
    for text in [
        "Plan: a planner reads the ledger (every experiment with its hypothesis, metrics and verdict) and proposes "
        "the next experiments inside a declared search space (four model families, three feature sets, class "
        "weighting, calibration). The Claude planner returns a structured object; code checks ranges, drops "
        "duplicates and falls back to the rules planner on any failure.",
        "Run: 5-fold stratified cross-validation on the train part with identical folds for every experiment; "
        "preprocessing lives inside the pipeline, so it is fit inside each fold.",
        "Judge: fold-paired comparison against the incumbent. Confirmed needs a mean gain of at least 0.002 PR-AUC "
        "and a paired t statistic of at least 2; smaller gains are marked inconclusive, refuted ideas stay in the trail.",
        "Narrate: a prompt turns the leaderboard, trail, holdout metrics, drivers and retention table into a "
        "conclusion with next experiments and retention actions; a template narrator produces the same sections "
        "without a model.",
        "The holdout is scored once, for the winner only, and reported next to the CV mean.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    rules = load_run("rules")
    claude = load_run("claude")
    doc.add_heading("Results", level=2)
    intro = "Primary metric PR-AUC (average precision); a random ranking scores 0.265, the churn rate."
    if rules and claude:
        intro += " Both planners ran with the same seed, folds, round limit and experiments per round."
    doc.add_paragraph(intro)
    header = ["planner", "experiments", "rounds", "confirmed gains", "winner", "CV PR-AUC (+- fold std)", "holdout PR-AUC", "holdout ROC-AUC", "lift@10%", "wall time"]
    rows = []
    for label, run in (("rules", rules), ("Claude", claude)):
        if run is None:
            continue
        f, r = run["final"], run["run"]
        rows.append([
            label, r["n_experiments"], r["n_rounds"], r["n_confirmed"], f["winner"]["name"],
            f"{f['winner_cv']['pr_auc']:.3f} (+- {f['winner_cv']['pr_auc_std']:.3f})",
            f"{f['holdout']['pr_auc']:.3f}", f"{f['holdout']['roc_auc']:.3f}", f"{f['holdout']['lift_at_10']:.2f}",
            f"{r['seconds']:.0f} s",
        ])
    add_table(doc, header, rows)
    if rules:
        f, r = rules["final"], rules["run"]
        doc.add_paragraph()
        doc.add_paragraph(
            "The finding on this dataset: no experiment beat the first logistic-regression baseline beyond fold noise; "
            "boosting at default settings overfits and regularising it closes only part of the gap. The rules run "
            f"stopped after round {r['n_rounds']} ({r['stop_reason']}). Holdout precision in the top decile is "
            f"{f['holdout']['precision_at_10']:.3f} against a base rate of {f['card']['churn_rate_holdout']:.3f}."
        )

    doc.add_heading("Screenshots", level=2)
    for name, caption in CAPTIONS.items():
        path = SHOTS / name
        if not path.exists():
            continue
        doc.add_picture(str(path), width=Cm(17))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(caption)
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].italic = True

    docx_path = OUT_DIR / "AI_challenge_submission_Aiman_El_Asad.docx"
    doc.save(docx_path)
    print("written", docx_path, f"({words}-word summary)")

    if not args.no_pdf:
        pdf_path = docx_path.with_suffix(".pdf")
        ps = (
            "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
            f"$d = $w.Documents.Open('{docx_path}'); $d.ExportAsFixedFormat('{pdf_path}', 17); "
            "$d.Close(0); $w.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
        print("written", pdf_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
